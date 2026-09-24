from pathlib import Path

import openpyxl

from scraper import Book, next_page_url, parse_page, write_output

FIX = Path(__file__).parent / "fixtures"
PAGE1 = (FIX / "page1.html").read_text(encoding="utf-8")
LAST = (FIX / "page50.html").read_text(encoding="utf-8")
BASE = "https://books.toscrape.com/catalogue/page-1.html"


def test_parse_page_extracts_20_books_with_fields():
    books = parse_page(PAGE1, BASE)
    assert len(books) == 20
    b = books[0]
    assert b.title == "A Light in the Attic"
    assert b.price == 51.77
    assert b.rating == 3
    assert b.in_stock is True
    assert b.url == "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html"


def test_pagination():
    assert next_page_url(PAGE1, BASE) == "https://books.toscrape.com/catalogue/page-2.html"
    assert next_page_url(LAST, "https://books.toscrape.com/catalogue/page-50.html") is None


def test_write_xlsx_and_csv(tmp_path):
    books = [Book("T", 1.5, 4, True, "http://x")]
    write_output(books, tmp_path / "out.xlsx")
    ws = openpyxl.load_workbook(tmp_path / "out.xlsx").active
    assert [c.value for c in ws[1]] == ["title", "price_gbp", "rating", "in_stock", "url"]
    assert [c.value for c in ws[2]] == ["T", 1.5, 4, True, "http://x"]
    write_output(books, tmp_path / "out.csv")
    assert (tmp_path / "out.csv").read_text().splitlines()[1] == "T,1.5,4,True,http://x"
